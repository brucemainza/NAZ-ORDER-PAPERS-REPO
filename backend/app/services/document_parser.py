"""Bounded and isolated parsing for uploaded parliamentary documents."""

from __future__ import annotations

from contextlib import contextmanager
from io import BytesIO
import multiprocessing
import os
from pathlib import Path, PurePosixPath
import tempfile
from typing import BinaryIO, Iterator, TypedDict
from zipfile import BadZipFile, ZipFile

try:
    import resource
except ImportError:  # pragma: no cover - resource is POSIX-only (no-op on Windows)
    resource = None  # type: ignore[assignment]

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - required in production images
    DocxDocument = None  # type: ignore[misc, assignment]

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - required in production images
    PdfReader = None  # type: ignore[misc, assignment]


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
UPLOAD_CHUNK_BYTES = 64 * 1024
MAX_SUBJECT_LENGTH = 500
MAX_EXTRACTED_CHARS = 2_000_000
MAX_PDF_PAGES = 200
MAX_ARCHIVE_ENTRIES = 500
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_COMPRESSION_RATIO = 200
DEFAULT_PARSE_TIMEOUT_SECONDS = 15.0
DEFAULT_PARSE_MEMORY_MB = 1024

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
EXPECTED_MIME_TYPES = {
    "pdf": {"application/pdf"},
    "docx": {DOCX_MIME},
    "txt": {"text/plain"},
}


class ParsedDocument(TypedDict):
    item_type: str | None
    subject: str
    full_text: str


class DocumentParsingTimeout(ValueError):
    pass


class DocumentParsingResourceLimit(ValueError):
    pass


def _detect_format(filename: str | None) -> str:
    if not filename:
        raise ValueError("Filename is required")
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in EXPECTED_MIME_TYPES:
        raise ValueError(
            f"Unsupported file type: '{extension}'. Only PDF, DOCX and TXT are supported."
        )
    return extension


def _validate_declared_mime(fmt: str, content_type: str | None) -> None:
    if content_type is None:
        return
    normalized = content_type.split(";", 1)[0].strip().casefold()
    if normalized not in EXPECTED_MIME_TYPES[fmt]:
        raise ValueError(
            f"Declared MIME type '{normalized or 'missing'}' does not match .{fmt}"
        )


def _validate_docx_archive(source: BinaryIO) -> None:
    source.seek(0)
    try:
        with ZipFile(source) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise ValueError("File content does not match the .docx extension")
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("DOCX exceeds safe archive limits")

            uncompressed_total = 0
            for entry in entries:
                member_path = PurePosixPath(entry.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise ValueError("DOCX contains an unsafe archive path")
                if entry.flag_bits & 0x1:
                    raise ValueError("Encrypted DOCX files are not supported")
                uncompressed_total += entry.file_size
                ratio = entry.file_size / max(1, entry.compress_size)
                if ratio > MAX_ARCHIVE_COMPRESSION_RATIO:
                    raise ValueError("DOCX exceeds safe archive limits")
            if uncompressed_total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX exceeds safe archive limits")
    except BadZipFile as error:
        raise ValueError("File content does not match the .docx extension") from error
    finally:
        source.seek(0)


def _validate_magic(source: BinaryIO, fmt: str) -> None:
    source.seek(0)
    header = source.read(8)
    source.seek(0)
    if fmt == "pdf" and not header.startswith(b"%PDF-"):
        raise ValueError("File content does not match the .pdf extension")
    if fmt == "docx":
        if not header.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            raise ValueError("File content does not match the .docx extension")
        _validate_docx_archive(source)
    if fmt == "txt":
        if b"\x00" in header:
            raise ValueError("File content does not match the .txt extension")
        try:
            source.read().decode("utf-8")
        except UnicodeDecodeError as error:
            raise ValueError("TXT file is not valid UTF-8") from error
        finally:
            source.seek(0)


def _extract_pdf_text(source: BinaryIO) -> str:
    if PdfReader is None:
        raise RuntimeError("PDF parser is not available")
    source.seek(0)
    try:
        reader = PdfReader(source, strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise ValueError(
                f"PDF has too many pages ({page_count}); maximum is {MAX_PDF_PAGES}"
            )
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Unable to read PDF file") from error

    parts: list[str] = []
    extracted_chars = 0
    for page in reader.pages:
        try:
            text = page.extract_text()
        except Exception as error:
            raise ValueError("Failed to extract text from PDF page") from error
        if text:
            extracted_chars += len(text)
            if extracted_chars > MAX_EXTRACTED_CHARS:
                raise ValueError("Extracted document text exceeds the safe limit")
            parts.append(text)
    return "\n".join(parts)


def _extract_docx_text(source: BinaryIO) -> str:
    if DocxDocument is None:
        raise RuntimeError("DOCX parser is not available")
    source.seek(0)
    try:
        document = DocxDocument(source)
    except Exception as error:
        raise ValueError("Unable to read DOCX file") from error

    parts: list[str] = []
    extracted_chars = 0
    for paragraph in document.paragraphs:
        if paragraph.text:
            extracted_chars += len(paragraph.text)
            if extracted_chars > MAX_EXTRACTED_CHARS:
                raise ValueError("Extracted document text exceeds the safe limit")
            parts.append(paragraph.text)
    return "\n".join(parts)


def _extract_txt_text(source: BinaryIO) -> str:
    source.seek(0)
    try:
        return source.read().decode("utf-8")
    except UnicodeDecodeError as error:
        raise ValueError("TXT file is not valid UTF-8") from error


def _clean_text(raw_text: str) -> str:
    paragraphs = raw_text.splitlines()
    cleaned_paragraphs = [" ".join(paragraph.split()) for paragraph in paragraphs]
    return "\n".join(line for line in cleaned_paragraphs if line)


def _split_subject_and_body(text: str) -> tuple[str, str]:
    lines = text.splitlines()
    if not lines:
        raise ValueError("Document contains no extractable text")
    subject = lines[0].strip()
    body = "\n\n".join(line.strip() for line in lines[1:] if line.strip())
    if not subject:
        raise ValueError("Document subject line is empty")
    return subject[:MAX_SUBJECT_LENGTH], body


def _parse_source(
    source: BinaryIO,
    *,
    filename: str | None,
    content_type: str | None,
    item_type: str | None,
) -> ParsedDocument:
    fmt = _detect_format(filename)
    _validate_declared_mime(fmt, content_type)
    _validate_magic(source, fmt)
    if fmt == "pdf":
        raw_text = _extract_pdf_text(source)
    elif fmt == "docx":
        raw_text = _extract_docx_text(source)
    else:
        raw_text = _extract_txt_text(source)

    text = _clean_text(raw_text)
    if not text:
        raise ValueError("No readable text found in the uploaded document")
    subject, body = _split_subject_and_body(text)
    return {
        "item_type": item_type if item_type in {"Question", "Motion"} else None,
        "subject": subject,
        "full_text": body,
    }


def parse_uploaded_document(
    content: bytes,
    filename: str | None,
    item_type: str | None = None,
    content_type: str | None = None,
) -> ParsedDocument:
    """Compatibility entry point for already-bounded in-memory content."""

    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError(
            f"File is too large ({len(content):,} bytes). "
            f"Maximum is {MAX_UPLOAD_BYTES:,} bytes."
        )
    return _parse_source(
        BytesIO(content),
        filename=filename,
        content_type=content_type,
        item_type=item_type,
    )


def parse_uploaded_document_path(
    path: str | Path,
    *,
    filename: str | None,
    content_type: str | None,
    item_type: str | None,
) -> ParsedDocument:
    with Path(path).open("rb") as source:
        return _parse_source(
            source,
            filename=filename,
            content_type=content_type,
            item_type=item_type,
        )


@contextmanager
def bounded_upload_file(
    source: BinaryIO,
    *,
    max_bytes: int = MAX_UPLOAD_BYTES,
    temp_directory: str | None = None,
) -> Iterator[Path]:
    """Copy an upload in bounded chunks and remove the temp file on every path."""

    descriptor, raw_path = tempfile.mkstemp(
        prefix="naz-upload-",
        suffix=".pending",
        dir=temp_directory,
    )
    path = Path(raw_path)
    total = 0
    try:
        with os.fdopen(descriptor, "wb") as destination:
            while True:
                chunk = source.read(UPLOAD_CHUNK_BYTES)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise ValueError(
                        f"File is too large. Maximum is {max_bytes:,} bytes."
                    )
                destination.write(chunk)
            destination.flush()
            os.fsync(destination.fileno())
        yield path
    finally:
        path.unlink(missing_ok=True)


def _apply_resource_limits(memory_mb: int, timeout_seconds: float) -> None:
    if resource is None:  # pragma: no cover - platform-dependent
        return
    memory_bytes = memory_mb * 1024 * 1024
    resource.setrlimit(resource.RLIMIT_AS, (memory_bytes, memory_bytes))
    cpu_seconds = max(1, int(timeout_seconds) + 1)
    resource.setrlimit(resource.RLIMIT_CPU, (cpu_seconds, cpu_seconds))


def _parse_worker(connection, kwargs: dict, memory_mb: int, timeout: float) -> None:
    try:
        _apply_resource_limits(memory_mb, timeout)
        result = parse_uploaded_document_path(**kwargs)
        connection.send(("ok", result))
    except Exception as error:
        connection.send(("error", error.__class__.__name__, str(error)))
    finally:
        connection.close()


def parse_document_safely(
    path: str | Path,
    *,
    filename: str | None,
    content_type: str | None,
    item_type: str | None,
    timeout_seconds: float = DEFAULT_PARSE_TIMEOUT_SECONDS,
    memory_mb: int = DEFAULT_PARSE_MEMORY_MB,
) -> ParsedDocument:
    """Parse in a killable process with wall-clock, CPU and memory limits."""

    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(
        target=_parse_worker,
        args=(
            child,
            {
                "path": str(path),
                "filename": filename,
                "content_type": content_type,
                "item_type": item_type,
            },
            memory_mb,
            timeout_seconds,
        ),
        daemon=True,
    )
    process.start()
    child.close()
    try:
        if not parent.poll(timeout_seconds):
            process.terminate()
            process.join(timeout=1)
            if process.is_alive():
                process.kill()
                process.join(timeout=1)
            raise DocumentParsingTimeout("Document parsing timed out")
        message = parent.recv()
    except EOFError as error:
        raise DocumentParsingResourceLimit(
            "Document parsing exceeded a resource limit"
        ) from error
    finally:
        parent.close()
        if process.is_alive():
            process.join(timeout=1)

    process.join(timeout=1)
    if message[0] == "ok":
        return message[1]
    error_name, detail = message[1], message[2]
    if error_name in {"MemoryError", "DocumentParsingResourceLimit"}:
        raise DocumentParsingResourceLimit(
            "Document parsing exceeded the memory limit"
        )
    if error_name == "RuntimeError":
        raise RuntimeError(detail)
    raise ValueError(detail)
